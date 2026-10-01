import sqlite3
from pathlib import Path
import openpyxl

BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = BASE_DIR / "baza.xlsx"
DB_FILE = BASE_DIR / "baza.db"

EMPLOYEE_SHEET = "Xodimlar"
STORE_SHEET = "Do'konlar bazasi"


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize_code(value):
    value = clean(value).upper()

    for symbol in [":", "-", "_", " "]:
        value = value.replace(symbol, ".")

    return value


print("=" * 50)
print("EXCEL → SQLITE")
print("=" * 50)

# Excel mavjudligini tekshirish
if not EXCEL_FILE.exists():
    print(f"❌ Excel topilmadi:")
    print(EXCEL_FILE)
    input("Enter bosing...")
    exit()

print("📂 Excel ochilmoqda...")

wb = openpyxl.load_workbook(
    EXCEL_FILE,
    data_only=True
)

# ============================================================
# DATABASE
# ============================================================

conn = sqlite3.connect(DB_FILE)
cursor = conn.cursor()

# SQLite tezligi uchun
cursor.execute("PRAGMA journal_mode=WAL")
cursor.execute("PRAGMA synchronous=NORMAL")

# ============================================================
# ESKI JADVALLARNI O'CHIRISH
# ============================================================

cursor.execute("DROP TABLE IF EXISTS employees")
cursor.execute("DROP TABLE IF EXISTS stores")
cursor.execute("DROP TABLE IF EXISTS requests")

# ============================================================
# XODIMLAR
# ============================================================

cursor.execute("""
CREATE TABLE employees (
    id INTEGER PRIMARY KEY,
    name TEXT NOT NULL,
    position TEXT,
    delivery_code TEXT,
    telegram_id INTEGER
)
""")

# ============================================================
# DO'KONLAR
# ============================================================

cursor.execute("""
CREATE TABLE stores (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    store_code TEXT UNIQUE NOT NULL,
    store_name TEXT,
    region TEXT,
    delivery_code TEXT
)
""")

# ============================================================
# SOROVLAR
# ============================================================

cursor.execute("""
CREATE TABLE requests (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    telegram_id INTEGER,
    employee_name TEXT,
    username TEXT,
    store_code TEXT,
    store_name TEXT,
    address TEXT,
    district TEXT,
    old_delivery TEXT,
    new_delivery TEXT,
    reason TEXT,
    status TEXT,
    manager_comment TEXT,
    operator_comment TEXT,
    manager_id INTEGER,
    manager_name TEXT,
    created_at TEXT,
    updated_at TEXT
)
""")

# ============================================================
# XODIMLAR
# ============================================================

print()
print("👥 Xodimlar ko'chirilmoqda...")

employee_count = 0

if EMPLOYEE_SHEET not in wb.sheetnames:
    print(f"❌ '{EMPLOYEE_SHEET}' sheet topilmadi!")
else:

    ws = wb[EMPLOYEE_SHEET]

    for row in ws.iter_rows(min_row=2, values_only=True):

        name = clean(row[0] if len(row) > 0 else "")
        employee_id = clean(row[1] if len(row) > 1 else "")
        position = clean(row[2] if len(row) > 2 else "")
        delivery_code = clean(row[3] if len(row) > 3 else "")

        if not employee_id:
            continue

        try:
            employee_id_int = int(employee_id)
        except (ValueError, TypeError):
            print(f"⚠️ ID o'tkazilmadi: {employee_id}")
            continue

        cursor.execute("""
        INSERT OR REPLACE INTO employees
        (
            id,
            name,
            position,
            delivery_code,
            telegram_id
        )
        VALUES (?, ?, ?, ?, ?)
        """, (
            employee_id_int,
            name,
            position,
            delivery_code,
            None
        ))

        employee_count += 1

# ============================================================
# DO'KONLAR
# ============================================================

print("🏪 Do'konlar ko'chirilmoqda...")

store_count = 0

if STORE_SHEET not in wb.sheetnames:
    print(f"❌ '{STORE_SHEET}' sheet topilmadi!")
else:

    ws = wb[STORE_SHEET]

    for row in ws.iter_rows(min_row=2, values_only=True):

        # A = Do'kon kodi
        store_code = clean(
            row[0] if len(row) > 0 else ""
        )

        # B = Do'kon nomi
        store_name = clean(
            row[1] if len(row) > 1 else ""
        )

        # C = Region
        region = clean(
            row[2] if len(row) > 2 else ""
        )

        # D = Biriktirilgan Dastavka
        delivery_code = clean(
            row[3] if len(row) > 3 else ""
        )

        if not store_code:
            continue

        normalized_code = normalize_code(store_code)

        cursor.execute("""
        INSERT OR REPLACE INTO stores
        (
            store_code,
            store_name,
            region,
            delivery_code
        )
        VALUES (?, ?, ?, ?)
        """, (
            normalized_code,
            store_name,
            region,
            delivery_code
        ))

        store_count += 1

# ============================================================
# INDEXLAR
# ============================================================

print("⚡ Qidiruv indekslari yaratilmoqda...")

cursor.execute("""
CREATE INDEX IF NOT EXISTS idx_store_code
ON stores(store_code)
""")

cursor.execute("""
CREATE INDEX IF NOT EXISTS idx_delivery_code
ON stores(delivery_code)
""")

cursor.execute("""
CREATE INDEX IF NOT EXISTS idx_employee_delivery
ON employees(delivery_code)
""")

cursor.execute("""
CREATE INDEX IF NOT EXISTS idx_request_status
ON requests(status)
""")

cursor.execute("""
CREATE INDEX IF NOT EXISTS idx_request_telegram
ON requests(telegram_id)
""")

# ============================================================
# SAQLASH
# ============================================================

conn.commit()

conn.close()

wb.close()

# ============================================================
# NATIJA
# ============================================================

print()
print("=" * 50)
print("✅ KO'CHIRISH TUGADI!")
print("=" * 50)

print(f"👥 Xodimlar: {employee_count} ta")
print(f"🏪 Do'konlar: {store_count} ta")
print(f"💾 SQL baza:")
print(DB_FILE)

print()
print("📁 SQL baza yaratildi:")
print("baza.db")
print()

input("Davom etish uchun Enter bosing...")